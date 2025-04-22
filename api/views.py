from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from django.utils.dateparse import parse_date
from openpyxl import load_workbook
from django.http import JsonResponse
from rest_framework import generics


import os
import mimetypes
from django.utils.encoding import smart_str
from django.http import FileResponse, Http404
from rest_framework.decorators import action


from rest_framework import viewsets, permissions
from rest_framework.permissions import AllowAny

from .models import (
    Patient,
    PatientDocument,
    Doctor,
    DoctorDocument,
    Service,
    ClinicSchedule,
    Appointment,
    License,
    Specialty, 
    User
)
from .serializers import (
    PatientSerializer,
    PatientDocumentSerializer,
    DoctorSerializer,
    DoctorDocumentSerializer,
    ServiceSerializer,
    ClinicScheduleSerializer,
    AppointmentSerializer,
    LicenseSerializer,
    SpecialtySerializer,
    UserSerializer
)





class RegisterView(generics.CreateAPIView):
    queryset = User.objects.all()
    serializer_class = UserSerializer
    permission_classes = [AllowAny]



class PatientViewSet(viewsets.ModelViewSet):
    """
    API endpoint для операций с пациентами.
    """
    queryset = Patient.objects.all()
    serializer_class = PatientSerializer



class PatientDocumentViewSet(viewsets.ModelViewSet):
    """
    API endpoint для операций с документами пациента.
    Фильтруем по текущему пациенту или по ?patient=...
    """
    serializer_class = PatientDocumentSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        qs = PatientDocument.objects.all()

        patient_id = self.request.query_params.get('patient')
        if patient_id:
            qs = qs.filter(patient__id=patient_id)
        else:
            user = self.request.user
            if hasattr(user, 'patient_profile'):
                qs = qs.filter(patient=user.patient_profile)

        return qs

    @action(detail=True, methods=['get'], url_path='download')
    def download(self, request, pk=None):
        doc = self.get_object()
        file_path = doc.document_file.path
        if not os.path.exists(file_path):
            raise Http404("Файл не найден")

        mime_type, _ = mimetypes.guess_type(file_path)
        mime_type = mime_type or 'application/octet-stream'

        response = FileResponse(open(file_path, 'rb'), content_type=mime_type)
        response['Content-Disposition'] = (
            f'attachment; filename="{smart_str(os.path.basename(file_path))}"'
        )
        return response

class DoctorViewSet(viewsets.ModelViewSet):
    queryset = Doctor.objects.all()
    serializer_class = DoctorSerializer
    http_method_names = ['get']

    def get_queryset(self):
        specialty = self.request.query_params.get('specialty')
        if specialty:
            return self.queryset.filter(specialty=specialty)
        return self.queryset


class DoctorDocumentViewSet(viewsets.ModelViewSet):
    """
    API endpoint для операций с документами врача.
    """
    queryset = DoctorDocument.objects.all()
    serializer_class = DoctorDocumentSerializer



class ServiceViewSet(viewsets.ModelViewSet):
    queryset = Service.objects.all()
    serializer_class = ServiceSerializer
    http_method_names = ['get']



class ClinicScheduleViewSet(viewsets.ModelViewSet):
    queryset = ClinicSchedule.objects.all()
    serializer_class = ClinicScheduleSerializer
    http_method_names = ['get']

    def get_queryset(self):
        doctor_id = self.request.query_params.get('doctor')
        weekday = self.request.query_params.get('weekday')
        queryset = self.queryset

        if doctor_id:
            queryset = queryset.filter(doctor__id=doctor_id)
        if weekday:
            queryset = queryset.filter(weekday=weekday)

        return queryset




class AppointmentViewSet(viewsets.ModelViewSet):
    """
    API endpoint для операций с приёмами.
    Теперь фильтруем по текущему пациенту или по ?patient=...
    """
    serializer_class = AppointmentSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        qs = Appointment.objects.filter(is_deleted=False)

        # 1) Фильтрация по patient query param
        patient_id = self.request.query_params.get('patient')
        if patient_id:
            qs = qs.filter(patient__id=patient_id)
        else:
            # 2) Если не передали query param, возьмём профиль текущего пользователя
            user = self.request.user
            if hasattr(user, 'patient_profile'):
                qs = qs.filter(patient=user.patient_profile)
            # иначе (врач/админ) — оставляем весь список

        # 3) Сохраняем старую логику фильтрации по врачу и дате
        doctor_id = self.request.query_params.get('doctor')
        if doctor_id:
            qs = qs.filter(doctor__id=doctor_id)

        appointment_date = self.request.query_params.get('appointment_date')
        if appointment_date:
            qs = qs.filter(appointment_date=appointment_date)

        return qs

    def perform_create(self, serializer):
        serializer.save()






class LicenseViewSet(viewsets.ModelViewSet):
    """
    API endpoint для операций с лицензиями клиники.
    """
    queryset = License.objects.all()
    serializer_class = LicenseSerializer


# Направление | Врач (ФИО)  | Кабинет | День недели | Время начала | Время конца | Дата (необяз.) | Длительность
# Кардиология | Иванов И.И. |   101   | Понедельник |     08:00    |    12:00    | (можно пусто)  | 15
# Терапия     | Петрова А.В.|   102   |   Вторник   |     13:00    |    17:00    |   2025-04-20   | 20


class ScheduleUploadView(APIView):
    def post(self, request):
        excel_file = request.FILES.get("file")

        if not excel_file:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = load_workbook(excel_file)
            sheet = wb.active

            for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True)):  # skip header
                direction, doctor_name, cabinet, weekday_str, start, end, specific_date, duration = row

                # Поиск врача по ФИО
                try:
                    last, first, patronymic = doctor_name.split()
                    doctor = Doctor.objects.get(last_name=last, first_name=first, patronymic=patronymic)
                except Exception:
                    return Response({"error": f"Doctor not found for row {i+2}"}, status=status.HTTP_400_BAD_REQUEST)

                weekday_map = {
                    'Понедельник': 1, 'Вторник': 2, 'Среда': 3,
                    'Четверг': 4, 'Пятница': 5, 'Суббота': 6, 'Воскресенье': 7
                }

                ClinicSchedule.objects.create(
                    direction=direction,
                    doctor=doctor,
                    cabinet=str(cabinet),
                    weekday=weekday_map.get(weekday_str, 1),
                    specific_date=parse_date(str(specific_date)) if specific_date else None,
                    start_time=start,
                    end_time=end,
                    appointment_duration=int(duration)
                )

            return Response({"status": "Расписание загружено успешно!"})
        
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
        


def get_doctors_by_specialty(request):
    specialty_id = request.GET.get('specialty_id')
    doctors = Doctor.objects.filter(specialty_id=specialty_id).values('id', 'first_name', 'last_name')
    return JsonResponse(list(doctors), safe=False)
    


class SpecialtyList(generics.ListAPIView):
    queryset = Specialty.objects.all()
    serializer_class = SpecialtySerializer

class ClinicScheduleList(generics.ListAPIView):
    queryset = ClinicSchedule.objects.all()
    serializer_class = ClinicScheduleSerializer